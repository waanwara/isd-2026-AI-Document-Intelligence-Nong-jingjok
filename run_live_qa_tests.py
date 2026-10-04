"""รันการทดสอบ 28 คำถามกับ HTTP Server จริง (http://127.0.0.1:8000/ask) แล้วบันทึกลง docs/tester.xlsx.

โหมดใช้งาน:
  python run_live_qa_tests.py                          # ยิงเซิร์ฟเวอร์สด (ต้องเปิด server ก่อน)
  python run_live_qa_tests.py --replay docs/reviewed_run.json
                                                       # ไม่ยิงเซิร์ฟเวอร์ ใช้คำตอบรอบที่คนตรวจไว้

คอลัมน์ในรายงาน:
  - "ตรวจคำสำคัญ (auto)": ตรวจแค่ว่าคำสำคัญปรากฏในคำตอบ เป็นเกณฑ์หลวม ไม่ใช่การยืนยันว่าถูก
  - "ตรวจโดยคน": อ่านเทียบเฉลยโดยคน เก็บใน docs/manual_review.json ผูกกับข้อความคำตอบ (sha1)
    ถ้ารันสดแล้วคำตอบเปลี่ยน ผลตรวจจะไม่ถูกนำมาใช้ และขึ้นว่า "ยังไม่ได้ตรวจ"
  - การตรวจ citation รายข้อไม่ได้ทำในสคริปต์นี้ ดู python -m katrag.eval.qa_eval
"""

import argparse
import hashlib
import json
import time
import urllib.request
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# List of 28 test items
TEST_ITEMS = [
    # EASY 1-9
    {
        "id": "E1", "level": "ระดับ 1: ง่าย", "program": "DSBA",
        "question": "ปีหนึ่งเทอมหนึ่งเรียนวิชาอะไร",
        "skill": "ค้นหาฟิลด์/ข้อความตรง ๆ จากเล่มหลักสูตร + อ้างอิงหน้า",
        "expected": "วิชาบังคับปี 1 เทอม 1 ของ DSBA 2565 มี 7 วิชา 18 หน่วยกิต: แคลคูลัส 1, พีชคณิตเชิงเส้น, พื้นฐานทางธุรกิจสำหรับเทคโนโลยีสารสนเทศ, การแก้ปัญหาและการโปรแกรมคอมพิวเตอร์, โรงเรียนสร้างเสน่ห์, กีฬาและนันทนาการ, ภาษาอังกฤษพื้นฐาน 1",
        "ref": "แผนการศึกษา DSBA 2565 (หน้า 14)",
    },
    {
        "id": "E2", "level": "ระดับ 1: ง่าย", "program": "DSBA",
        "question": "วิชาแคลคูลัส 2 มีกี่หน่วยกิต",
        "skill": "ค้นหาฟิลด์/ข้อความตรง ๆ จากเล่มหลักสูตร",
        "expected": "3 หน่วยกิต [3(3-0-6)]",
        "ref": "ตารางรายวิชา DSBA 2565 (หน้า 12)",
    },
    {
        "id": "E3", "level": "ระดับ 1: ง่าย", "program": "DSBA",
        "question": "วิชาการสร้างคลังข้อมูลเรียนปีไหน เทอมอะไร",
        "skill": "ค้นหาฟิลด์/ข้อความตรง ๆ จากเล่มหลักสูตร",
        "expected": "ปีที่ 3 ภาคการศึกษาที่ 1",
        "ref": "แผนการศึกษา DSBA 2565 (หน้า 15)",
    },
    {
        "id": "E4", "level": "ระดับ 1: ง่าย", "program": "IT",
        "question": "ปีสองเทอมสองเรียนวิชาอะไร",
        "skill": "ค้นหาฟิลด์/ข้อความตรง ๆ จากเล่มหลักสูตร",
        "expected": "รายวิชาบังคับ IT ปี 2 เทอม 2 ตามแผนการศึกษา (Cybersecurity, Software Engineering, OS, NoSQL, Functional Prog, Web Prog ฯลฯ)",
        "ref": "แผนการศึกษา IT 2565",
    },
    {
        "id": "E5", "level": "ระดับ 1: ง่าย", "program": "DSBA",
        "question": "วิชาระบบข้อมูลมหัต ชื่อภาษาอังกฤษว่าอะไร",
        "skill": "ค้นหาฟิลด์/ข้อความตรง ๆ จากเล่มหลักสูตร",
        "expected": "BIG DATA SYSTEMS",
        "ref": "ตารางรายวิชา DSBA 2565",
    },
    {
        "id": "E6", "level": "ระดับ 1: ง่าย", "program": "AIT",
        "question": "ปีหนึ่งเรียนอะไรบ้าง",
        "skill": "ค้นหาฟิลด์/ข้อความตรง ๆ จากเล่มหลักสูตร",
        "expected": "รายวิชาบังคับ AIT ปี 1 ทั้ง 2 ภาคการศึกษา (เทอม 1: Calculus 1, Linear Alg, Discrete, Prob/Stat, PSCP / เทอม 2: Calculus 2, Comp Prog, Embedded, Data Struct, Team Project 1, Digital Citizen)",
        "ref": "แผนการศึกษา AIT 2566",
    },
    {
        "id": "E7", "level": "ระดับ 1: ง่าย", "program": "DSBA",
        "question": "วิชาแคลคูลัส 2 ต้องผ่านวิชาใดก่อน",
        "skill": "ค้นหาข้อความ Prerequisite ตรง ๆ",
        "expected": "แคลคูลัส 1 (06026200)",
        "ref": "ตารางรายวิชา DSBA 2565 (หน้า 12)",
    },
    {
        "id": "E8", "level": "ระดับ 1: ง่าย", "program": "DSBA",
        "question": "วิชาการเขียนโปรแกรมมีกี่หน่วยกิต?",
        "skill": "ค้นหาฟิลด์/ข้อความตรง ๆ จากเล่มหลักสูตร (คำถามจากสไลด์ P2)",
        "expected": "วิชา 06026203 การโปรแกรมคอมพิวเตอร์ มี 3 หน่วยกิต [3(2-2-5)] และวิชา 06066303 การแก้ปัญหาและการโปรแกรมคอมพิวเตอร์ มี 3 หน่วยกิต [3(2-2-5)]",
        "ref": "ตารางรายวิชา DSBA 2565 (หน้า 12)",
    },
    {
        "id": "E9", "level": "ระดับ 1: ง่าย", "program": "IT",
        "question": "รหัสวิชา 06066303 ชื่อวิชาอะไร?",
        "skill": "ค้นหาชื่อวิชาจากรหัสวิชาตรง ๆ (Code Lookup จากสไลด์ P2)",
        "expected": "การแก้ปัญหาและการโปรแกรมคอมพิวเตอร์ (PROBLEM SOLVING AND COMPUTER PROGRAMMING)",
        "ref": "ตารางรายวิชาหมวดวิชาเฉพาะกลุ่มพื้นฐานวิชาชีพ",
    },

    # MEDIUM 1-9
    {
        "id": "M1", "level": "ระดับ 2: ปานกลาง", "program": "DSBA",
        "question": "วิชา data warehouse ต้องผ่านวิชาใดก่อน และวิชานั้นเรียนตอนไหน",
        "skill": "อ่านตาราง + รวม / กรองข้อมูลหลายจุด + อ้างอิงหน้า",
        "expected": "ต้องผ่าน 06066300 แนวคิดระบบฐานข้อมูล (DATABASE SYSTEM CONCEPTS) ซึ่งเรียนปีที่ 2 ภาคการศึกษาที่ 1",
        "ref": "Prerequisite + แผนการศึกษา DSBA 2565 (หน้า 15, 18)",
    },
    {
        "id": "M2", "level": "ระดับ 2: ปานกลาง", "program": "DSBA",
        "question": "มีวิชาเขียนโปรแกรมกี่หน่วยกิต",
        "skill": "อ่านตาราง + รวม / กรองข้อมูลหลายจุด",
        "expected": "4 วิชา รวม 12 หน่วยกิต (PSCP, Computer Programming, Data Analytics & Programming, Fundamental Web Programming)",
        "ref": "ตารางรายวิชา DSBA 2565 (หน้า 12–15)",
    },
    {
        "id": "M3", "level": "ระดับ 2: ปานกลาง", "program": "DSBA",
        "question": "วิชาคณิตศาสตร์มีกี่ตัว",
        "skill": "อ่านตาราง + รวม / กรองข้อมูลหลายจุด",
        "expected": "6 วิชา รวม 18 หน่วยกิต (Calculus 1, Calculus 2, Linear Algebra, Discrete Math, Probability & Statistics, Bayesian Statistics)",
        "ref": "ตารางรายวิชา DSBA 2565",
    },
    {
        "id": "M4", "level": "ระดับ 2: ปานกลาง", "program": "IT",
        "question": "วิชาเกี่ยวกับเครือข่ายมีกี่วิชา",
        "skill": "อ่านตาราง + รวม / กรองข้อมูลหลายจุด",
        "expected": "7 วิชา รวม 21 หน่วยกิต (Intro to Networks, Communication Network Infrastructure, IoT, Wireless Network, Network Design, Network Performance, Troubleshooting)",
        "ref": "ตารางรายวิชา IT 2565",
    },
    {
        "id": "M5", "level": "ระดับ 2: ปานกลาง", "program": "AIT",
        "question": "วิชาเกี่ยวกับปัญญาประดิษฐ์มีอะไรบ้าง",
        "skill": "อ่านตาราง + รวม / กรองข้อมูลหลายจุด",
        "expected": "14 วิชา รวม 42 หน่วยกิต เช่น AI & IoT, Seminar in AI, Project in AI 1-2, AI Ethics, AI Service Design, Selected Topics in AI 1-6",
        "ref": "ตารางรายวิชา AIT 2566",
    },
    {
        "id": "M6", "level": "ระดับ 2: ปานกลาง", "program": "DSBA",
        "question": "ปีสามเรียนอะไร",
        "skill": "อ่านตาราง + รวม / กรองข้อมูลหลายจุด",
        "expected": "วิชาบังคับปี 3 (เทอม 1: Applied ML, Data Warehousing, Professional Communication / เทอม 2: Big Data Systems, โครงงาน 1, IT Project Management, ผู้ประกอบการสมัยใหม่) + ต้องเลือกวิชาแขนง",
        "ref": "แผนการศึกษา DSBA 2565 (หน้า 27–29)",
    },
    {
        "id": "M7", "level": "ระดับ 2: ปานกลาง", "program": "DSBA",
        "question": "วิชาบังคับชั้นปี 2 ภาคต้นมีอะไรบ้าง?",
        "skill": "อ่านตาราง + กรองข้อมูลเฉพาะภาคเรียน (คำถามจากสไลด์ P2)",
        "expected": "มี 5 วิชาบังคับ 15 หน่วยกิต ได้แก่ 1) การวิเคราะห์ข้อมูลและการโปรแกรม 2) คณิตศาสตร์ไม่ต่อเนื่อง 3) แนวคิดระบบฐานข้อมูล 4) การเขียนโปรแกรมเว็บพื้นฐาน 5) การวิเคราะห์และออกแบบระบบสารสนเทศ",
        "ref": "แผนการศึกษา DSBA 2565 (หน้า 14–15)",
    },
    {
        "id": "M8", "level": "ระดับ 2: ปานกลาง", "program": "DSBA",
        "question": "หมวดวิชาเฉพาะเลือกเก็บกี่หน่วยกิต?",
        "skill": "อ่านตาราง + รวม / กรองข้อมูลหมวดวิชา (คำถามจากสไลด์ P2)",
        "expected": "ไม่น้อยกว่า 12 หน่วยกิต (กลุ่มวิชาชีพเลือกด้าน 3 แขนง) และมีกลุ่มวิชาการศึกษาทางเลือกอีก 6 หน่วยกิต",
        "ref": "โครงสร้างหลักสูตร DSBA 2565 (หน้า 15)",
    },
    {
        "id": "M9", "level": "ระดับ 2: ปานกลาง", "program": "DSBA",
        "question": "วิชา Data Warehousing ต้องผ่านวิชาใดก่อน? ถ้ายังไม่ผ่านจะลงทะเบียนได้ไหม",
        "skill": "เชื่อมโยงเงื่อนไข prerequisite + ให้เหตุผล (คำถามจากสไลด์ P2)",
        "expected": "ต้องสอบผ่าน 06066300 แนวคิดระบบฐานข้อมูล ก่อน หากยังไม่ผ่าน จะไม่สามารถลงทะเบียนเรียนได้",
        "ref": "คำอธิบายรายวิชา DSBA 2565 (หน้า 18)",
    },

    # HARD 1-7
    {
        "id": "H1", "level": "ระดับ 3: ยาก", "program": "DSBA",
        "question": "ถ้าจะจบใน 3.5 ปี แต่ละเทอมต้องลงวิชาอะไร",
        "skill": "เชื่อมโยงเงื่อนไข prerequisite + ให้เหตุผล + วางแผน (สไลด์ P2)",
        "expected": "ต้องดึงวิชาที่ปกติอยู่ปี 4 เทอม 2 (+ วิชาเลือก) มากระจายลงเทอมก่อนหน้า เฉลี่ยเพิ่มเทอมละ 3-6 หน่วยกิต และต้องตรวจ prerequisite",
        "ref": "แผนการศึกษา DSBA 2565 (หน้า 14–16)",
    },
    {
        "id": "H2", "level": "ระดับ 3: ยาก", "program": "DSBA",
        "question": "เรียนวิชา database ปีสี่ได้ไหม",
        "skill": "เชื่อมโยงเงื่อนไข + ให้เหตุผลและตีความระเบียบ",
        "expected": "ได้ — เอกสารระบุ 'แผนการศึกษา' ที่แนะนำ ไม่ใช่ข้อห้าม และไม่มีข้อความห้ามลงวิชาบังคับช้ากว่าแผน และยังมีวิชาฐานข้อมูลที่เปิดปี 4 จริง คือ 06026246 ระบบฐานข้อมูลแบบกระจาย (ปี 4 เทอม 1 ไม่มีวิชาบังคับก่อน)",
        "ref": "แผนการศึกษา + คำอธิบายรายวิชา (หน้า 15, 323)",
    },
    {
        "id": "H3", "level": "ระดับ 3: ยาก", "program": "DSBA",
        "question": "ลงวิชา data warehouse ตอนปีสองได้ไหม เพราะอะไร",
        "skill": "เชื่อมโยงเงื่อนไข prerequisite + ให้เหตุผล",
        "expected": "ไม่ได้ตามแผนปกติ — วิชาอยู่ปี 3 เทอม 1 และต้องผ่านแนวคิดระบบฐานข้อมูล (ปี 2 เทอม 1) ก่อน",
        "ref": "Prerequisite + แผนการศึกษา DSBA 2565",
    },
    {
        "id": "H4", "level": "ระดับ 3: ยาก", "program": "DSBA",
        "question": "วิชาที่มีในหลักสูตรเก่า แต่ไม่มีในหลักสูตรใหม่ มีอะไรบ้าง",
        "skill": "เชื่อมโยงข้อมูลข้ามเวอร์ชัน + Set Difference (สไลด์ P2)",
        "expected": "วิชาที่อยู่ใน DSBA 2560 แต่ไม่มีชื่อตรงใน 2565 มี 48 วิชา เช่น เทคโนโลยีเว็บ (06026109), การวิเคราะห์ข้อมูลเชิงธุรกิจ (06026116), ระบบฐานข้อมูลขั้นสูง (06026145)",
        "ref": "เทียบตารางรายวิชา DSBA 2560 vs 2565",
    },
    {
        "id": "H5", "level": "ระดับ 3: ยาก", "program": "DSBA",
        "question": "หมวดวิชาเฉพาะเลือกเก็บกี่หน่วยกิต",
        "skill": "เชื่อมโยงเงื่อนไข + แยกแยะหมวดวิชา",
        "expected": "ไม่น้อยกว่า 12 หน่วยกิต — โครงสร้างหลักสูตรระบุ 'กลุ่มวิชาชีพเลือกด้าน' 12 หน่วยกิต ส่วนเลข 6 หน่วยกิตคือกลุ่มวิชาการศึกษาทางเลือก ซึ่งเป็นคนละหมวด",
        "ref": "โครงสร้างหลักสูตร DSBA 2565 (หน้า 15)",
    },
    {
        "id": "H6", "level": "ระดับ 3: ยาก", "program": "DSBA",
        "question": "ปีสามต้องเลือกแขนงไหม มีแขนงอะไรให้เลือก",
        "skill": "เชื่อมโยงเงื่อนไข + สรุปกลุ่มความเชี่ยวชาญ",
        "expected": "ต้องเลือกแขนง — มี 3 กลุ่ม: วิทยาการข้อมูล, การวิเคราะห์เชิงสถิติ, วิศวกรรมข้อมูล",
        "ref": "วิชาเลือกกลุ่มความเชี่ยวชาญ DSBA 2565",
    },
    {
        "id": "H7", "level": "ระดับ 3: ยาก", "program": "DSBA",
        "question": "ตรวจว่าแผนเรียนนี้ครบเงื่อนไขจบหรือไม่: นักศึกษา DSBA เก็บศึกษาทั่วไป 30 หน่วยกิต, เฉพาะบังคับ 78 หน่วยกิต, เลือกเสรี 6 หน่วยกิต แต่เก็บเฉพาะเลือกได้ 9 หน่วยกิต รวม 123 หน่วยกิต",
        "skill": "เชื่อมโยงเงื่อนไขจบ + ตรวจสอบกฎระเบียบ (สไลด์ P2)",
        "expected": "ยังไม่ครบเงื่อนไขสำเร็จการศึกษา ขาดวิชาเฉพาะเลือก 3 หน่วยกิต (เกณฑ์ 12 หน่วยกิต) และหน่วยกิตรวมได้ 123 หน่วยกิต ขาดจากเกณฑ์จบขั้นต่ำ 132 หน่วยกิต อยู่ 9 หน่วยกิต",
        "ref": "เกณฑ์การสำเร็จการศึกษา DSBA 2565 (หน้า 7)",
    },

    # CHALLENGE 1-3 (ระดับ 4 ท้าทาย +10 คะแนนพิเศษ)
    {
        "id": "C1", "level": "ระดับ 4: ท้าทาย", "program": "IT",
        "question": "เปรียบเทียบหลักสูตร IT 2560 (เล่มเก่า) กับ IT 2565 (เล่มปัจจุบัน) มีการปรับโครงสร้างหน่วยกิตรวมและแขนงวิชาอย่างไรบ้าง?",
        "skill": "Multi-hop ข้ามเวอร์ชันหลักสูตร + วางแผน + สังเคราะห์ (สไลด์ P2)",
        "expected": "หน่วยกิตรวมปรับลดจาก 130 หน่วยกิต เหลือ 129 หน่วยกิต และจัดโครงสร้างแขนงวิชาใหม่ชัดเจนเป็น 3 แขนง (Software Engineering, Network and System Technology, Multimedia and Game Development)",
        "ref": "IT2560_old.pdf (หน้า 10–18) เปรียบเทียบกับ IT2565_current.pdf (หน้า 15–25, 372)",
    },
    {
        "id": "C2", "level": "ระดับ 4: ท้าทาย", "program": "AITBA",
        "question": "ในเอกสารหลักสูตรระดับบัณฑิตศึกษา (ป.โท และ ป.เอก) มีคู่หลักสูตรใดที่ใช้เอกสารเล่มเดียวกัน และมีข้อสังเกตอย่างไร?",
        "skill": "Multi-hop + วิเคราะห์ความสัมพันธ์เอกสารข้ามระดับ (สไลด์ P2)",
        "expected": "ไฟล์ PH_D_AITBA2569_current.pdf (ป.เอก) และ M_AITBA2569_current.pdf (ป.โท) เป็นเอกสารเล่มเดียวกันทุกประการ (ขนาด 12,285,167 ไบต์, SHA256 เท่ากัน, จำนวน 252 หน้า) ระบุหน้าปกเป็นหลักสูตรวิทยาศาสตรมหาบัณฑิต (วท.ม.) ทั้งคู่",
        "ref": "เอกสารหลักสูตร AITBA 2569 หน้า 1–2 และ readme.txt",
    },
    {
        "id": "C3", "level": "ระดับ 4: ท้าทาย", "program": "DSBA",
        "question": "เปรียบเทียบหลักสูตร DSBA 2565 แผนปกติ (ไม่ทำสหกิจ) กับ แผนสหกิจศึกษา มีความแตกต่างกันในภาคเรียนใด และมีวิชาใดที่ต่างกัน?",
        "skill": "Multi-hop + แตกแขนงแผนการเรียน (Academic Plan Branching)",
        "expected": "ต่างกันใน ปี 4 เทอม 2 โดยแผนปกติจะลงเรียนโครงงาน 2 (3 หน่วยกิต) ควบคู่กับวิชาเลือก ส่วนแผนสหกิจศึกษาจะลงวิชาสหกิจศึกษา 6 หน่วยกิตไปปฏิบัติงานเต็มเวลา",
        "ref": "แผนการศึกษาเปรียบเทียบ DSBA 2565 (หน้า 15–17)",
    },
]

# ─────────────────────────────────────────────────────────────────────────────
# Run tests
# ─────────────────────────────────────────────────────────────────────────────
ap = argparse.ArgumentParser(description="Run the 28-question benchmark and write docs/tester.xlsx")
ap.add_argument("--replay", help="ใช้คำตอบจากไฟล์ JSON ที่เก็บไว้ แทนการยิงเซิร์ฟเวอร์สด")
ap.add_argument("--url", default="http://127.0.0.1:8000/ask")
ARGS = ap.parse_args()

REVIEW_PATH = Path("docs/manual_review.json")
REVIEWS = {}
REVIEW_META = {}
if REVIEW_PATH.exists():
    _rv = json.loads(REVIEW_PATH.read_text(encoding="utf-8"))
    REVIEWS = _rv.get("reviews", {})
    REVIEW_META = _rv.get("_meta", {})

DOC_MAP = {
    "a39d6db17ccc5cf4": "DSBA 2565",
    "344608973458106b": "DSBA 2560",
    "5dec80d93328c9fa": "IT 2565",
    "09c745bf0cd1fefe": "IT 2560",
    "71905b5244a14b94": "AIT 2566",
    "75085b1dd7523d89": "BIT 2565",
    "1ffe70b5db234caa": "BIT 2560",
    "bef8aad4da2cad3d": "ปร.ด. AITBA 2569",
    "430d1625db23e79b": "วท.ม. AITBA 2569",
    "a2bb37ed5f089453": "วท.ม. AITBA 2564",
    "272a77249680d878": "วท.ม. IT 2568",
    "0bb1dc8496421928": "วท.ม. IT 2563",
    "184978cecf5b14e7": "ปร.ด. IT 2566",
    "4d03ccad7edd646c": "ปร.ด. IT 2561",
}


def format_doc(doc_id):
    if doc_id in DOC_MAP:
        return DOC_MAP[doc_id]
    if str(doc_id).endswith(".pdf"):
        return str(doc_id).replace(".pdf", "").split("/")[-1]
    return str(doc_id)


def fetch_live(item):
    payload = {"question": item["question"], "program": item["program"]}
    req = urllib.request.Request(
        ARGS.url, data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    t0 = time.time()
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            body = json.loads(resp.read().decode("utf-8"))
        return {"answer": body.get("answer", ""), "citations": body.get("citations", []), "latency": time.time() - t0, "error": None}
    except Exception as e:  # noqa: BLE001
        return {"answer": f"Error: {e}", "citations": [], "latency": time.time() - t0, "error": str(e)}


replay = {}
if ARGS.replay:
    replay = {r["id"]: r for r in json.loads(Path(ARGS.replay).read_text(encoding="utf-8"))}
    print(f"Replaying {len(replay)} saved results from {ARGS.replay} (no server calls)\n")
else:
    print(f"Starting live test run for {len(TEST_ITEMS)} items against {ARGS.url} ...\n")


def check_accuracy(item_id: str, ans: str) -> str:
    """ตรวจแบบคำสำคัญ (หลวม): คืน PASS ถ้าคำสำคัญทุกคำปรากฏในคำตอบ"""
    if "Error:" in ans or not ans:
        return "FAIL"
    ans_l = ans.lower()
    checks = {
        "E1": ["แคลคูลัส", "06026200"], "E2": ["3"], "E3": ["ปีที่ 3", "3"], "E4": ["ปีที่ 2", "ซอฟต์แวร์"],
        "E5": ["big data systems"], "E6": ["แคลคูลัส", "06046400"], "E7": ["แคลคูลัส 1", "06026200"], "E8": ["3"],
        "E9": ["การแก้ปัญหา", "problem solving"], "M1": ["แนวคิดระบบฐานข้อมูล", "06066300"], "M2": ["12"], "M3": ["18"],
        "M4": ["21"], "M5": ["42"], "M6": ["ปีที่ 3", "การสร้างคลังข้อมูล"], "M7": ["15"], "M8": ["12"],
        "M9": ["ไม่ได้", "06066300"], "H1": ["3.5"], "H2": ["ได้"], "H3": ["ไม่ได้"], "H4": ["2560", "2565"],
        "H5": ["12"], "H6": ["3", "วิทยาการข้อมูล"], "H7": ["ยังไม่ครบ", "ขาด"], "C1": ["129", "130"],
        "C2": ["m_aitba2569", "ph_d_aitba2569"], "C3": ["ปีที่ 4", "สหกิจ"],
    }
    for req in checks.get(item_id, []):
        if req.lower() not in ans_l:
            return "FAIL"
    return "PASS"


live_results = []
for item in TEST_ITEMS:
    qid = item["id"]
    if ARGS.replay:
        r = replay[qid]
        got = {"answer": r["answer"], "citations": r.get("citations", []), "latency": r["latency"], "error": None}
    else:
        got = fetch_live(item)
    answer, citations, elapsed = got["answer"], got["citations"], got["latency"]
    print(f"[{qid}] {elapsed:.2f}s | citations: {len(citations)}" + (f" | ERROR: {got['error']}" if got["error"] else ""))

    grouped_pages = {}
    for c in citations:
        dname = format_doc(c.get("document_id", ""))
        p = c.get("page", "")
        grouped_pages.setdefault(dname, [])
        if p and str(p) not in grouped_pages[dname]:
            grouped_pages[dname].append(str(p))
    ref_txt = item["ref"]
    if grouped_pages:
        live_cite_str = "; ".join(f"{d} หน้า {', '.join(pages[:4])}" for d, pages in grouped_pages.items())
        ref_txt = f"{item['ref']}\n(หน้าที่ระบบคืนมา: {live_cite_str})"

    rv = REVIEWS.get(qid, {})
    sha = hashlib.sha1(answer.encode("utf-8")).hexdigest()
    if rv and rv.get("answer_sha1") == sha:
        verdict, note = rv["verdict"], rv["note"]
    elif rv:
        verdict, note = "ยังไม่ได้ตรวจ", "คำตอบรอบนี้ต่างจากรอบที่คนตรวจ (LLM ตอบไม่เหมือนเดิม) ต้องอ่านเทียบเฉลยใหม่"
    else:
        verdict, note = "ยังไม่ได้ตรวจ", ""

    live_results.append({
        "id": qid, "level": item["level"], "program": item["program"], "question": item["question"],
        "skill": item["skill"], "expected": item["expected"], "answer": answer, "ref": ref_txt,
        "path": rv.get("path", "ไม่ทราบ (ยังไม่ได้ trace)"), "path_kind": rv.get("path_kind", ""),
        "n_cite": len(citations), "kw": check_accuracy(qid, answer), "verdict": verdict, "note": note,
        "latency": f"{elapsed:.2f}s", "latency_val": elapsed, "latency_pass": "PASS" if elapsed < 5.0 else "FAIL",
    })

with open("docs/live_results_cache.json", "w", encoding="utf-8") as f:
    json.dump(live_results, f, ensure_ascii=False, indent=2)

print("\nAll queries finished. Generating docs/tester.xlsx ...")

# ─────────────────────────────────────────────────────────────────────────────
# Workbook
# ─────────────────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
font_title = Font(name="Segoe UI", size=16, bold=True, color="1E293B")
font_subtitle = Font(name="Segoe UI", size=11, italic=True, color="64748B")
font_section = Font(name="Segoe UI", size=13, bold=True, color="FF6600")
font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
font_cell = Font(name="Segoe UI", size=9, color="1E293B")
font_bold = Font(name="Segoe UI", size=9, bold=True, color="1E293B")
font_pass = Font(name="Segoe UI", size=9, bold=True, color="15803D")
font_warn = Font(name="Segoe UI", size=9, bold=True, color="B45309")
font_bad = Font(name="Segoe UI", size=9, bold=True, color="B91C1C")
fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
fill_sub = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
fill_easy = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
fill_medium = PatternFill(start_color="FEFCE8", end_color="FEFCE8", fill_type="solid")
fill_hard = PatternFill(start_color="FFF7ED", end_color="FFF7ED", fill_type="solid")
fill_challenge = PatternFill(start_color="FAF5FF", end_color="FAF5FF", fill_type="solid")
fill_pass = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
fill_warn = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
fill_bad = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
side = Side(border_style="thin", color="CBD5E1")
thin_border = Border(left=side, right=side, top=side, bottom=side)
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
align_center_top = Alignment(horizontal="center", vertical="top", wrap_text=True)

LEVELS = [("ง่าย", "ระดับ 1: ง่าย"), ("ปานกลาง", "ระดับ 2: ปานกลาง"), ("ยาก", "ระดับ 3: ยาก"), ("ท้าทาย", "ระดับ 4: ท้าทาย")]


def by_level(key):
    return [r for r in live_results if key in r["level"]]


def verdict_style(cell, v):
    cell.alignment = align_center_top
    if v == "ถูกครบ":
        cell.fill, cell.font = fill_pass, font_pass
    elif v == "ถูกบางส่วน":
        cell.fill, cell.font = fill_warn, font_warn
    elif v == "ผิด":
        cell.fill, cell.font = fill_bad, font_bad
    else:
        cell.fill, cell.font = fill_sub, font_bold


n_total = len(live_results)
n_kw = sum(1 for r in live_results if r["kw"] == "PASS")
n_full = sum(1 for r in live_results if r["verdict"] == "ถูกครบ")
n_part = sum(1 for r in live_results if r["verdict"] == "ถูกบางส่วน")
n_wrong = sum(1 for r in live_results if r["verdict"] == "ผิด")
n_unrev = sum(1 for r in live_results if r["verdict"] == "ยังไม่ได้ตรวจ")
times = [r["latency_val"] for r in live_results]
n_hard = [r["id"] for r in live_results if r["path_kind"] == "hardcoded"]

# ── Sheet 1: Summary ──
ws = wb.active
ws.title = "Executive Summary"
ws["A1"] = "Jingjok-Thorius — รายงานผลทดสอบระบบถาม-ตอบหลักสูตร"
ws["A1"].font = font_title
ws["A2"] = "ผลยิงทดสอบผ่าน HTTP endpoint POST /ask (ตัวเลขทั้งหมดคำนวณจากข้อมูลในแผ่น All Test Cases)"
ws["A2"].font = font_subtitle

meta = [
    ("ระบบเป้าหมาย:", "Jingjok-Thorius HTTP server (http://127.0.0.1:8000)"),
    ("สถาปัตยกรรม:", "Structured SQL (SQLite) + Hybrid retrieval (FTS5 + bge-m3) + Typhoon LLM สำหรับคำถามเชิงเหตุผล"),
    ("LLM:", "Typhoon v2.5 30B-A3B (typhoon-v2.5-30b-a3b-instruct) ผ่าน OpenAI-compatible API"),
    ("ข้อมูล:", "14 เล่มหลักสูตร (3,689 หน้า, 1,419 รายวิชา)"),
    ("จำนวนข้อทดสอบ:", f"{n_total} ข้อ (E1–E9, M1–M9, H1–H7, C1–C3)"),
    ("ตรวจคำสำคัญ (auto):", f"{n_kw}/{n_total} ข้อ — เกณฑ์หลวม (บางข้อตรวจแค่มีตัวเลข) ไม่ใช่การยืนยันว่าถูก"),
    ("ตรวจโดยคน (เทียบเฉลย):", f"ถูกครบ {n_full} · ถูกบางส่วน {n_part} · ผิด {n_wrong} · ยังไม่ได้ตรวจ {n_unrev}"),
    ("เวลาตอบ:", f"เฉลี่ย {sum(times)/n_total:.2f}s · เร็วสุด {min(times):.2f}s · ช้าสุด {max(times):.2f}s · ต่ำกว่า 5s {sum(t < 5 for t in times)}/{n_total} ข้อ"),
    ("วิธีตรวจโดยคน:", REVIEW_META.get("method", "-") + f" · ตรวจเมื่อ {REVIEW_META.get('reviewed_on', '-')}"),
]
for idx, (k, v) in enumerate(meta, start=4):
    ws.cell(row=idx, column=1, value=k).font = font_bold
    c = ws.cell(row=idx, column=2, value=v)
    c.font = font_cell
    ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=8)

start_row = 15
ws.cell(row=start_row - 1, column=1, value="สรุปแยกตามระดับความยาก").font = font_section
hdr = ["ระดับ", "จำนวนข้อ", "ตรวจคำสำคัญผ่าน", "ถูกครบ (คนตรวจ)", "ถูกบางส่วน", "ผิด", "เวลาเฉลี่ย (s)", "เวลาช้าสุด (s)"]
for ci, h in enumerate(hdr, start=1):
    c = ws.cell(row=start_row, column=ci, value=h)
    c.font, c.fill, c.alignment = font_header, fill_header, align_center
rows_out = []
for key, label in LEVELS:
    g = by_level(key)
    rows_out.append((label, len(g), sum(r["kw"] == "PASS" for r in g), sum(r["verdict"] == "ถูกครบ" for r in g),
                     sum(r["verdict"] == "ถูกบางส่วน" for r in g), sum(r["verdict"] == "ผิด" for r in g),
                     f"{sum(r['latency_val'] for r in g)/len(g):.2f}", f"{max(r['latency_val'] for r in g):.2f}"))
rows_out.append(("รวม", n_total, n_kw, n_full, n_part, n_wrong, f"{sum(times)/n_total:.2f}", f"{max(times):.2f}"))
for ro, row in enumerate(rows_out, start=1):
    for ci, val in enumerate(row, start=1):
        c = ws.cell(row=start_row + ro, column=ci, value=val)
        c.font = font_bold if (ci == 1 or ro == len(rows_out)) else font_cell
        c.border = thin_border
        c.alignment = align_left if ci == 1 else align_center
        if ro == len(rows_out):
            c.fill = fill_sub

nr = start_row + len(rows_out) + 2
ws.cell(row=nr, column=1, value="ข้อควรทราบเมื่ออ่านผล").font = font_section
notes = [
    "1) 'ตรวจคำสำคัญ' ผ่านง่าย (เช่น E2 ตรวจแค่มีเลข 3) จึงไม่ควรใช้เป็นอัตราความถูกต้อง ให้ดูคอลัมน์ 'ตรวจโดยคน'",
    "2) ผลตรวจโดยคนอ่านเทียบเฉลยในไฟล์นี้ ไม่ได้เทียบกับ PDF ต้นฉบับโดยตรง และผูกกับคำตอบรอบที่ตรวจ (คำถามเชิงเหตุผลเรียก LLM คำตอบเปลี่ยนได้ทุกรอบ)",
    "3) ความถูกต้องของ citation ไม่ได้ตรวจรายข้อในไฟล์นี้ วัดแยกด้วย python -m katrag.eval.qa_eval (19 ข้อ): precision 0.57, recall 0.53",
    f"4) ข้อ {', '.join(n_hard)} ตอบด้วยข้อความสำเร็จรูปที่เขียนไว้ในโค้ด (structured_query.py) ไม่ได้ดึงจากฐานข้อมูลตอนถาม จึงไม่ควรใช้เป็นหลักฐานว่าระบบสกัดข้อมูลเหล่านั้นได้เอง",
    "5) คอลัมน์ 'เส้นทางประมวลผล' ได้จากการ trace โค้ด pipeline (นับว่ามีการเรียก LLM หรือไม่) ไม่ใช่ข้อมูลที่ API ส่งกลับมา",
    "6) เฉลยข้อ H4 เดิมเขียน 48 วิชา แต่ฐานข้อมูลปัจจุบันให้ 52 (DSBA 2560 71 ชื่อ / 2565 75 ชื่อ) ผลของระบบตรงกับฐานข้อมูล",
]
for i, t in enumerate(notes, start=1):
    c = ws.cell(row=nr + i, column=1, value=t)
    c.font = font_cell
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=nr + i, start_column=1, end_row=nr + i, end_column=8)
    ws.row_dimensions[nr + i].height = 30
for i, w in enumerate([30, 14, 18, 18, 14, 10, 16, 16], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ── Sheet 2: Test cases ──
wt = wb.create_sheet(title="All Test Cases & Results")
headers = [
    "ข้อที่", "ระดับ", "หลักสูตร", "คำถามทดสอบ", "ทักษะที่วัด", "เฉลย (Expected)", "คำตอบจริงจากเซิร์ฟเวอร์",
    "หน้าอ้างอิงที่ระบบคืนมา", "เส้นทางประมวลผล (trace จากโค้ด)", "จำนวน citation",
    "ตรวจคำสำคัญ (auto)", "ตรวจโดยคน (เทียบเฉลย)", "หมายเหตุผู้ตรวจ", "เวลาตอบ", "< 5s?",
]
for ci, h in enumerate(headers, start=1):
    c = wt.cell(row=1, column=ci, value=h)
    c.font, c.fill, c.alignment, c.border = font_header, fill_header, align_center, thin_border
wt.row_dimensions[1].height = 32
for ri, r in enumerate(live_results, start=2):
    wt.row_dimensions[ri].height = 80
    lvl_fill = fill_easy if "ง่าย" in r["level"] else fill_medium if "ปานกลาง" in r["level"] else fill_hard if "ยาก" in r["level"] else fill_challenge
    vals = [r["id"], r["level"], r["program"], r["question"], r["skill"], r["expected"], r["answer"], r["ref"], r["path"],
            r["n_cite"], r["kw"], r["verdict"], r["note"], r["latency"], r["latency_pass"]]
    for ci, v in enumerate(vals, start=1):
        c = wt.cell(row=ri, column=ci, value=v)
        c.font, c.border = font_cell, thin_border
        c.alignment = align_center_top if ci in (1, 2, 3, 10, 11, 12, 14, 15) else align_left_top
        if ci == 1:
            c.font = font_bold
        elif ci == 2:
            c.fill, c.font = lvl_fill, font_bold
        elif ci == 9 and r["path_kind"] == "hardcoded":
            c.fill, c.font = fill_warn, font_warn
        elif ci == 11 and v == "PASS":
            c.fill, c.font = fill_sub, font_bold
        elif ci == 12:
            verdict_style(c, v)
        elif ci == 15 and v == "PASS":
            c.fill, c.font = fill_pass, font_pass
        elif ri % 2 == 1 and ci not in (2, 11, 12, 15):
            c.fill = fill_zebra
for i, w in enumerate([8, 16, 10, 36, 28, 42, 52, 30, 34, 10, 12, 14, 44, 10, 8], start=1):
    wt.column_dimensions[get_column_letter(i)].width = w
wt.freeze_panes = "A2"
wt.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n_total + 1}"

# ── Sheet 3: Criteria (measured, no scores awarded) ──
wc = wb.create_sheet(title="Challenge Criteria (Measured)")
wc["A1"] = "เกณฑ์ Challenge — ผลที่วัดได้ (ไม่ได้ให้คะแนน คะแนนเป็นของผู้ตรวจ)"
wc["A1"].font = font_title
ch_hdr = ["เกณฑ์", "ข้อกำหนด", "ผลที่วัดได้", "ข้อสังเกต"]
for ci, h in enumerate(ch_hdr, start=1):
    c = wc.cell(row=3, column=ci, value=h)
    c.font, c.fill, c.alignment, c.border = font_header, fill_header, align_center, thin_border
chal = by_level("ท้าทาย")
c_lines = "; ".join(f"{r['id']}: {r['verdict']}" for r in chal)
hard_t = [r["latency_val"] for r in by_level("ยาก")]
crit_rows = [
    ("1. ระดับ 4 ท้าทาย (Multi-hop & Cross-version)",
     "สกัดเล่มหลายเวอร์ชัน เชื่อมโยงข้อมูลหลายจุด และอ้างอิงหน้า/ข้อในเล่มได้ถูกต้อง",
     c_lines,
     "C1 และ C3 ตอบด้วยข้อความสำเร็จรูปในโค้ด; C2 ดึงจากตาราง document_relation แต่ป้าย ป.โท/ป.เอก สลับ; C1 ส่วนแขนงวิชายังไม่ได้ยืนยันกับ PDF"),
    ("2. ความเร็ว (Latency)",
     "ตอบ 1 คำถามระดับยากภายใน < 5 วินาที (รวม retrieval + generate)",
     f"ระดับยาก {len(hard_t)} ข้อ เฉลี่ย {sum(hard_t)/len(hard_t):.2f}s ช้าสุด {max(hard_t):.2f}s · ทั้ง {n_total} ข้อ ต่ำกว่า 5s {sum(t < 5 for t in times)}/{n_total}",
     "วัดบนเซิร์ฟเวอร์ที่โหลดโมเดลแล้ว (warm); การถามครั้งแรกหลังเปิดเซิร์ฟเวอร์แบบ KATRAG_SKIP_WARMUP=1 จะช้ากว่า (โหลด bge-m3) ราว 0.5–1 นาที"),
]
for ri, row in enumerate(crit_rows, start=4):
    wc.row_dimensions[ri].height = 80
    for ci, v in enumerate(row, start=1):
        c = wc.cell(row=ri, column=ci, value=v)
        c.font = font_bold if ci == 1 else font_cell
        c.border = thin_border
        c.alignment = align_left_top
for i, w in enumerate([34, 46, 52, 60], start=1):
    wc.column_dimensions[get_column_letter(i)].width = w

output_path = Path("docs/tester.xlsx")
alt_path = Path("docs/tester_live.xlsx")
try:
    wb.save(output_path)
    print(f"\n[DONE] Saved results to {output_path}")
except PermissionError:
    print("\n[WARNING] docs/tester.xlsx is locked by Excel!")
try:
    wb.save(alt_path)
    print(f"[DONE] Saved copy to {alt_path}")
except Exception as e:  # noqa: BLE001
    print(f"[WARNING] Could not save to {alt_path}: {e}")
